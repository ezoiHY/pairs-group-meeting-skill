"""Deck structure for build_deck.py, made with python-pptx alone.

Each spec slide is a clone of one of the template's sample slides: a copy of the
sample's XML (background, title, logos) with its image relationships re-made on
the clone. The samples are then dropped from the slide list, and python-pptx
leaves out every part nothing refers to any more when it saves. The finished
deck is checked by reopening it.

This is the Codex version. The Claude version of the skill has its own
deck_structure.py with the same three functions, built on the pptx skill.
"""

import random
import zipfile
from copy import deepcopy

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
P14_CREATION_ID = "{http://schemas.microsoft.com/office/powerpoint/2010/main}creationId"


def add_arguments(ap):
    """No options: this version needs nothing outside the skill."""


def rel_refs(element):
    """(element, attribute, rId) for every relationship reference in an XML tree."""
    for el in element.iter("*"):
        for attr, value in el.attrib.items():
            if attr.startswith(R_NS):
                yield el, attr, value


def clone_slide(prs, src):
    dst = prs.slides.add_slide(src.slide_layout)
    new_rid = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        target = rel.target_ref if rel.is_external else rel.target_part
        new_rid[rid] = dst.part.relate_to(target, rel.reltype, is_external=rel.is_external)

    sld = dst._element
    for child in list(sld):
        sld.remove(child)
    for child in src._element:
        sld.append(deepcopy(child))
    # One pass, so an rId that is both an old and a new id is not renamed twice.
    for el, attr, rid in list(rel_refs(sld)):
        if rid in new_rid:
            el.set(attr, new_rid[rid])
    for el in sld.iter(P14_CREATION_ID):
        el.set("val", str(random.randint(1, 2**31 - 1)))
    return dst


def build(args, template, sources, originals, out_path):
    """Write a deck made of the template slides named in `sources`, in that order, without `originals`."""
    prs = Presentation(str(template))
    samples = {slide.part.partname.filename: slide for slide in prs.slides}
    sample_ids = list(prs.slides._sldIdLst)
    for name in sources:
        clone_slide(prs, samples[name])
    for sld_id in sample_ids:
        if prs.part.related_part(sld_id.rId).partname.filename in originals:
            prs.slides._sldIdLst.remove(sld_id)
            prs.part.drop_rel(sld_id.rId)
    prs.save(str(out_path))


def validate(args, out_path, template):
    """(ok, report): the deck reopens, and every relationship a slide refers to exists."""
    with zipfile.ZipFile(out_path) as zf:
        bad = zf.testzip()
        names = zf.namelist()
    if bad:
        return False, f"Check FAILED: corrupt zip entry {bad}"
    if len(names) != len(set(names)):
        return False, "Check FAILED: duplicate zip entries"
    prs = Presentation(str(out_path))
    problems = []
    for i, slide in enumerate(prs.slides, 1):
        for _, _, rid in rel_refs(slide._element):
            # A video's play action carries an empty r:id, as in PowerPoint's own files.
            if rid and rid not in slide.part.rels:
                problems.append(f"slide {i} refers to missing relationship {rid}")
    if problems:
        return False, "Check FAILED:\n" + "\n".join(problems)
    return True, f"Check passed: the deck reopens cleanly ({len(prs.slides)} slides)."
