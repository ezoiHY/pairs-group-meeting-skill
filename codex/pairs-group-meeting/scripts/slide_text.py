#!/usr/bin/env python3
"""Print the text of every slide in one or more decks, in slide order.

Used to check a built deck's content, and to read the previous meeting's deck
(its Todo slide is this meeting's Done candidate list).

Usage:
    python slide_text.py deck.pptx [more.pptx ...]
"""

import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def shape_lines(shape):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            yield from shape_lines(child)
    elif shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            text = "".join(r.text for r in p.runs).strip()
            if text:
                yield text
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            yield " | ".join(cell.text.strip() for cell in row.cells)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    for path in sys.argv[1:]:
        prs = Presentation(path)
        if len(sys.argv) > 2:
            print(f"=== {path}")
        for i, slide in enumerate(prs.slides, 1):
            print(f"--- slide {i}")
            for shape in slide.shapes:
                for line in shape_lines(shape):
                    print(line)


if __name__ == "__main__":
    main()
