#!/usr/bin/env python3
"""Build a PAIRS Lab group-meeting deck from a JSON spec.

The deck's structure is made with the pptx skill's own scripts: add_slide.py
clones the template's sample slides in spec order and clean.py drops the unused
originals. python-pptx then fills each clone, so every slide keeps the
template's background, logos and title styling. The finished deck is checked
with the pptx skill's validate.py against the template.

Usage:
    python build_deck.py spec.json [-o out.pptx] [--pptx-skill DIR] [--no-validate]

The spec format is documented in references/spec-format.md.
"""

import argparse
import datetime
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu

SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATE = SKILL_DIR / "assets" / "template.pptx"
CONFIG = SKILL_DIR / "config.json"

# Template sample slide each spec type is cloned from.
BASE_SLIDE = {"cover": "slide1.xml", "done": "slide2.xml", "content": "slide3.xml", "todo": "slide4.xml"}
TEMPLATE_SLIDES = ("slide1.xml", "slide2.xml", "slide3.xml", "slide4.xml")

# Geometry in EMU, measured from the template.
EMU_PER_PT = 12700
EMU_PER_IN = 914400
SLIDE_W, SLIDE_H = 12192000, 6858000
MARGIN_X = 349250                      # left edge of title and body text in the template
CONTENT_TOP = 1219200                  # top of the template's body text box
CONTENT_BOTTOM = 5740000               # 0.2" above the logo group, whose top is at 5922049
LOGO_LEFT = 8128524
FOOTNOTE_TOP, FOOTNOTE_BOTTOM = 5960000, 6400000
TITLE_W, TITLE_MAX_W = 7620000, 11000000
COVER_TITLE_W = 6096000
LIST_W = 8890000                       # width of the template's Done/Todo text box
FULL_W = SLIDE_W - 2 * MARGIN_X
GAP = int(0.2 * EMU_PER_IN)
COL_GAP = int(0.35 * EMU_PER_IN)
WIDE_FIGURE = 2.6                      # width/height from which a figure is larger under the text than beside it

# Colors from the template.
NAVY = "002147"        # body text
NAVY_DARK = "111F4A"   # cover title, table header
SLATE = "4A5878"       # cover subtitle, captions, footnotes
ONGOING = "980000"
DONE_SCHEME = "accent6"
RULE = "D5DBE6"
STRIPE = "F2F5FA"

FONTS_XML = (
    '<a:latin typeface="Calibri" panose="020F0502020204030204" pitchFamily="34" charset="0"/>'
    '<a:ea typeface="等线" panose="02010600030101010101" pitchFamily="2" charset="-122"/>'
    '<a:cs typeface="Calibri" panose="020F0502020204030204" pitchFamily="34" charset="0"/>'
)

DONE_WORDS = {"done", "finished", "complete", "completed", "完成", "已完成"}
ONGOING_WORDS = {"on-going", "ongoing", "in progress", "in-progress", "wip", "进行中"}

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}
CONVERT_EXTS = {".webp"}
RENDER_EXTS = {".pdf", ".svg"}
VIDEO_MIME = {".mp4": "video/mp4", ".m4v": "video/x-m4v", ".mov": "video/quicktime",
              ".avi": "video/x-msvideo", ".wmv": "video/x-ms-wmv"}

WARNINGS = []

# House style; config.json "style" overrides these.
STYLE = {"language": "en", "min_font_size": 22, "max_bullets_per_slide": 3, "max_words_per_bullet": 14,
         "max_compare_papers": 5, "max_compare_columns": 5}


def warn(msg):
    WARNINGS.append(msg)


def font_sizes(top):
    """Sizes to try for one kind of text: from `top` down to the house minimum, in 2pt steps."""
    floor = STYLE["min_font_size"]
    return list(range(max(top, floor), floor, -2)) + [floor]


# ---------------------------------------------------------------------------
# Text measurement

CJK = "\u1100-\u11ff\u2e80-\u9fff\uac00-\ud7af\uf900-\ufaff\ufe30-\ufe4f\uff00-\uffef"
CJK_RE = re.compile(f"[{CJK}]")
SEGMENT_RE = re.compile(f"[{CJK}]+|[^{CJK}]+")
TOKEN_RE = re.compile(f"[{CJK}]|[^\\s{CJK}]+\\s*|\\s+")

FONT_DIRS = [
    Path("/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts"),
    Path("/Library/Fonts"),
    Path.home() / "Library" / "Fonts",
    Path("C:/Windows/Fonts"),
    Path("/usr/share/fonts/truetype/crosextra"),
]
FONT_FILES = {
    "latin": ("Calibri.ttf", "calibri.ttf", "Carlito-Regular.ttf"),
    "latin_b": ("Calibrib.ttf", "calibrib.ttf", "Carlito-Bold.ttf"),
    "cjk": ("Deng.ttf", "deng.ttf"),
    "cjk_b": ("Dengb.ttf", "dengb.ttf", "Deng.ttf"),
}


class Metrics:
    """Text widths in points, from the deck's own fonts when they are installed."""

    def __init__(self):
        self.fonts = {}
        for key, names in FONT_FILES.items():
            for folder in FONT_DIRS:
                hit = next((folder / n for n in names if (folder / n).exists()), None)
                if hit:
                    self.fonts[key] = ImageFont.truetype(str(hit), 1000)
                    break

    def width(self, text, size, bold=False):
        total = 0.0
        for seg in SEGMENT_RE.findall(text):
            cjk = bool(CJK_RE.match(seg))
            font = self.fonts.get(("cjk" if cjk else "latin") + ("_b" if bold else ""))
            if font:
                total += font.getlength(seg) / 1000 * size
            else:
                total += sum(_guess_em(c, cjk) for c in seg) * size * (1.04 if bold else 1.0)
        return total


def _guess_em(c, cjk):
    if cjk:
        return 1.0
    if c.isspace():
        return 0.23
    if c.isupper():
        return 0.6
    if c.isalnum():
        return 0.5
    return 0.4


METRICS = None


def line_count(runs, width_emu, size):
    """Greedy word wrap, the way PowerPoint breaks lines."""
    limit = width_emu / EMU_PER_PT * 0.97
    lines, x = 1, 0.0
    for run in runs:
        for tok in TOKEN_RE.findall(run.text):
            w = METRICS.width(tok, size, run.bold)
            ink = METRICS.width(tok.rstrip(), size, run.bold)
            if x > 0 and x + ink > limit:
                lines += 1
                x = 0.0
                if tok.isspace():
                    continue
            if ink > limit:
                lines += int((x + ink) // limit)
                x = (x + w) % limit
                continue
            x += w
    return lines


def line_height(runs, size):
    # Measured in PowerPoint renders: single spacing is ~1.2x the size for Calibri and DengXian alike.
    has_cjk = any(CJK_RE.search(r.text) for r in runs)
    return size * (1.24 if has_cjk else 1.2)


# ---------------------------------------------------------------------------
# Paragraph model -> DrawingML


class Run:
    def __init__(self, text, bold=False, color=NAVY, scheme=None, italic=False):
        self.text, self.bold, self.color, self.scheme, self.italic = text, bold, color, scheme, italic


class Para:
    """One <a:p>. Spacing is given as a fraction of the font size so it scales with fitting."""

    def __init__(self, runs, marL=0, indent=0, before=0.0, bullet=None, algn="l", color=NAVY):
        self.runs, self.marL, self.indent = runs, marL, indent
        self.before, self.bullet, self.algn, self.color = before, bullet, algn, color

    def height(self, width_emu, size):
        if not self.runs:
            return line_height([], size) * EMU_PER_PT
        lines = line_count(self.runs, width_emu - self.marL, size)
        return (lines * line_height(self.runs, size) + self.before * size) * EMU_PER_PT

    def xml(self, size):
        if self.bullet:
            bu = (f'<a:buClr><a:srgbClr val="{self.color}"/></a:buClr><a:buSzPct val="100000"/>'
                  f'<a:buFont typeface="Arial" panose="020B0604020202020204" pitchFamily="34" charset="0"/>'
                  f'<a:buChar char="{escape(self.bullet)}"/>')
        else:
            bu = "<a:buNone/>"
        ppr = (f'<a:pPr marL="{self.marL}" indent="{self.indent}" algn="{self.algn}">'
               f'<a:lnSpc><a:spcPct val="100000"/></a:lnSpc>'
               f'<a:spcBef><a:spcPts val="{round(self.before * size * 100)}"/></a:spcBef>'
               f'<a:spcAft><a:spcPts val="0"/></a:spcAft>{bu}</a:pPr>')
        runs = "".join(run_xml(r, size) for r in self.runs)
        end = (f'<a:endParaRPr lang="en-US" altLang="zh-CN" sz="{round(size * 100)}" dirty="0">'
               f'<a:solidFill><a:srgbClr val="{self.color}"/></a:solidFill>{FONTS_XML}</a:endParaRPr>')
        return f'<a:p {nsdecls("a")}>{ppr}{runs}{end}</a:p>'


def run_xml(run, size):
    if run.scheme:
        fill = f'<a:solidFill><a:schemeClr val="{run.scheme}"/></a:solidFill>'
    else:
        fill = f'<a:solidFill><a:srgbClr val="{run.color}"/></a:solidFill>'
    attrs = f'lang="en-US" altLang="zh-CN" sz="{round(size * 100)}" b="{1 if run.bold else 0}"'
    if run.italic:
        attrs += ' i="1"'
    return f'<a:r><a:rPr {attrs} dirty="0">{fill}{FONTS_XML}</a:rPr><a:t>{escape(run.text)}</a:t></a:r>'


def inline_runs(text, bold=False, color=NAVY):
    """Split **bold** markup into runs."""
    runs = []
    for i, part in enumerate(str(text).split("**")):
        if part:
            runs.append(Run(part, bold=bold or i % 2 == 1, color=color))
    return runs


def stack_height(paras, width_emu, size):
    return sum(p.height(width_emu, size) for p in paras)


def write_paragraphs(tx_body, paras, size):
    for p in tx_body.findall(qn("a:p")):
        tx_body.remove(p)
    for p in paras:
        tx_body.append(parse_xml(p.xml(size)))


def add_textbox(slide, box, name, anchor="t"):
    x, y, w, h = (int(v) for v in box)
    shape = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    shape.name = name
    body_pr = shape.text_frame._txBody.find(qn("a:bodyPr"))
    for child in list(body_pr):
        body_pr.remove(child)
    for k, v in {"wrap": "square", "lIns": "0", "tIns": "0", "rIns": "0", "bIns": "0", "anchor": anchor}.items():
        body_pr.set(k, v)
    body_pr.append(parse_xml(f'<a:normAutofit {nsdecls("a")}/>'))
    return shape


def fit_size(paras, width_emu, height_emu, sizes):
    for size in sizes:
        if stack_height(paras, width_emu, size) <= height_emu:
            return size, True
    return sizes[-1], False


def set_single_text(shape, text, size):
    """Replace a placeholder's text but keep the template's first run formatting."""
    tx_body = shape.text_frame._txBody
    paras = tx_body.findall(qn("a:p"))
    for extra in paras[1:]:
        tx_body.remove(extra)
    p = paras[0]
    runs = p.findall(qn("a:r"))
    for extra in runs[1:]:
        p.remove(extra)
    runs[0].find(qn("a:t")).text = text
    for rpr in (runs[0].find(qn("a:rPr")), p.find(qn("a:endParaRPr"))):
        if rpr is not None:
            if size:
                rpr.set("sz", str(round(size * 100)))
            ensure_cjk_font(rpr)


def ensure_cjk_font(rpr):
    """Template runs name only a Latin font; give CJK text DengXian instead of the theme's SimSun."""
    if rpr.find(qn("a:ea")) is not None:
        return
    ea = parse_xml(f'<a:ea {nsdecls("a")} typeface="等线" panose="02010600030101010101" pitchFamily="2" charset="-122"/>')
    latin = rpr.find(qn("a:latin"))
    if latin is not None:
        latin.addnext(ea)
        return
    later = [rpr.find(qn(t)) for t in ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")]
    later = [e for e in later if e is not None]
    if later:
        later[0].addprevious(ea)
    else:
        rpr.append(ea)


# ---------------------------------------------------------------------------
# Slide fillers


def set_title(slide, text):
    shape = slide.shapes.title
    runs = [Run(text, bold=True)]
    width, size = TITLE_W, 36
    if line_count(runs, width, size) > 1:
        width = TITLE_MAX_W
        while size > 24 and line_count(runs, width, size) > 1:
            size -= 2
        if line_count(runs, width, size) > 1:
            warn(f"title too long for one line even at 24pt: {text!r}")
    shape.width = Emu(width)
    set_single_text(shape, text, size)


def fill_cover(slide, s):
    title = s.get("title") or "Group Meeting"
    shape = slide.shapes.title
    size = 48
    while size > 32 and line_count([Run(title, bold=True)], COVER_TITLE_W, size) > 1:
        size -= 2
    if line_count([Run(title, bold=True)], COVER_TITLE_W, size) > 1:
        shape.width = Emu(COVER_TITLE_W + 1000000)
    set_single_text(shape, title, size)

    sub = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 1)
    lines = [line for line in [s["date"], s.get("name", "")] + list(s.get("extra_lines", [])) if line]
    tx_body = sub.text_frame._txBody
    paras = tx_body.findall(qn("a:p"))
    while len(paras) < len(lines):
        clone = deepcopy(paras[-1])
        tx_body.append(clone)
        paras.append(clone)
    for extra in paras[len(lines):]:
        tx_body.remove(extra)
    for p, text in zip(paras, lines):
        runs = p.findall(qn("a:r"))
        for extra in runs[1:]:
            p.remove(extra)
        runs[0].find(qn("a:t")).text = text
        ensure_cjk_font(runs[0].find(qn("a:rPr")))
    if len(lines) > 2:
        sub.height = Emu(sub.height + (len(lines) - 2) * int(0.42 * EMU_PER_IN))


def normalize_item(item):
    if isinstance(item, str):
        return {"text": item}
    return item


def status_run(status):
    key = status.strip().lower()
    if key in DONE_WORDS:
        return Run("done" if key in {"done", "finished", "complete", "completed"} else status, scheme=DONE_SCHEME)
    if key in ONGOING_WORDS:
        return Run("on-going" if key in {"on-going", "ongoing", "wip", "in-progress"} else status, color=ONGOING)
    return Run(status, color=ONGOING)


def progress_paras(projects, kind):
    """Paragraphs of a Done/Todo list, styled like the template's sample text."""
    paras = []
    for pi, proj in enumerate(projects):
        label = proj.get("label", "Project: ")
        head = [Run(label, bold=True)] if label else []
        head += inline_runs(proj.get("name", ""), bold=True)
        paras.append(Para(head, before=0 if pi == 0 else 18 / 28))
        for ii, item in enumerate(proj.get("items", [])):
            item = normalize_item(item)
            runs = [Run(f"{chr(ord('a') + ii % 26)}. ")] + inline_runs(item["text"])
            if kind == "done" and item.get("status"):
                runs += [Run(" — "), status_run(item["status"])]
            paras.append(Para(runs, marL=457200, before=3 / 28))
    return paras


def plan_progress(s, kind):
    """Pick font size and width for a Done/Todo slide; split it when it cannot fit at the minimum size."""
    height = CONTENT_BOTTOM - CONTENT_TOP
    paras = progress_paras(s["projects"], kind)
    for size in font_sizes(28):
        for width in (LIST_W, FULL_W):
            if stack_height(paras, width, size) <= height:
                return [dict(s, _size=size, _width=width)]
    split_size = STYLE["min_font_size"]

    chunks, current = [], []
    for proj in s["projects"]:
        items = [normalize_item(i) for i in proj.get("items", [])]
        start = 0
        while True:
            name = proj.get("name", "") + (" (cont.)" if start else "")
            taken = []
            for item in items[start:]:
                trial = current + [dict(proj, name=name, items=taken + [item])]
                if stack_height(progress_paras(trial, kind), FULL_W, split_size) > height:
                    break
                taken.append(item)
            if not taken and items[start:] and current:
                chunks.append(current)       # retry this project on a fresh slide
                current = []
                continue
            if not taken and items[start:]:
                taken = [items[start]]       # a single item taller than a slide: keep it anyway
            current.append(dict(proj, name=name, items=taken))
            start += len(taken)
            if start >= len(items):
                break
            chunks.append(current)
            current = []
    if current:
        chunks.append(current)
    warn(f"'{s['title']}' did not fit on one slide; split into {len(chunks)} slides at {split_size}pt")
    out = []
    for i, chunk in enumerate(chunks):
        title = s["title"] if i == 0 else f"{s['title']} (cont.)"
        out.append(dict(s, title=title, projects=chunk, _size=split_size, _width=FULL_W))
    return out


def fill_progress(slide, s, kind):
    set_title(slide, s["title"])
    body = next(sh for sh in slide.shapes if sh.has_text_frame and not sh.is_placeholder)
    body.name = "Done list" if kind == "done" else "Todo list"
    body.left, body.top = Emu(MARGIN_X), Emu(CONTENT_TOP)
    body.width, body.height = Emu(s["_width"]), Emu(CONTENT_BOTTOM - CONTENT_TOP)
    write_paragraphs(body.text_frame._txBody, progress_paras(s["projects"], kind), s["_size"])


def bullet_paras(items, color=NAVY):
    paras = []
    for i, item in enumerate(items):
        item = normalize_item(item)
        level = int(item.get("level", 0))
        plain = item.get("bullet") is False
        runs = inline_runs(item["text"], bold=item.get("bold", False), color=item.get("color", color))
        before = 0 if i == 0 else (0.15 if level else 0.35)
        if plain:
            paras.append(Para(runs, marL=level * 342900, before=before, color=color))
        elif level:
            paras.append(Para(runs, marL=628650, indent=-228600, before=before, bullet="–", color=color))
        else:
            paras.append(Para(runs, marL=285750, indent=-285750, before=before, bullet="•", color=color))
    return paras


def place_text(slide, paras, box, sizes, name, what):
    x, y, w, h = box
    size, ok = fit_size(paras, w, h, sizes)
    if not ok:
        warn(f"{what}: text overflows its area even at {size}pt — cut words or split the slide")
    used = min(stack_height(paras, w, size), h) if ok else h
    shape = add_textbox(slide, (x, y, w, used + 0.1 * EMU_PER_IN), name)
    write_paragraphs(shape.text_frame._txBody, paras, size)
    return used, size


# ---------------------------------------------------------------------------
# Media


def prepare_media(item, base_dir, work):
    item = {"path": item} if isinstance(item, str) else dict(item)
    if "chart" in item:
        item.update(kind="chart", aspect=float(item.get("aspect", 1.6)))
        return item
    path = Path(os.path.expanduser(item["path"]))
    if not path.is_absolute():
        path = (base_dir / path).resolve()
    if not path.exists():
        raise SystemExit(f"media not found: {path}")
    ext = path.suffix.lower()
    item["src"] = path
    if ext in VIDEO_MIME:
        poster = item.get("poster")
        poster = Path(os.path.expanduser(poster)) if poster else video_poster(path, work)
        if not poster.is_absolute():
            poster = (base_dir / poster).resolve()
        with Image.open(poster) as im:
            item["aspect"] = im.width / im.height
        item.update(kind="video", poster=poster, mime=VIDEO_MIME[ext])
        return item
    if ext in RENDER_EXTS:
        import pymupdf
        doc = pymupdf.open(path)
        page = doc[int(item.get("page", 1)) - 1]
        out = work / f"{path.stem}-{len(list(work.iterdir()))}.png"
        page.get_pixmap(dpi=220, alpha=False).save(out)
        path = out
    elif ext in CONVERT_EXTS or ext not in IMAGE_EXTS:
        out = work / f"{path.stem}-{len(list(work.iterdir()))}.png"
        with Image.open(path) as im:
            im.save(out)
        path = out
    with Image.open(path) as im:
        upright = ImageOps.exif_transpose(im)
        if upright is not im:
            out = work / f"{path.stem}-upright-{len(list(work.iterdir()))}.png"
            upright.save(out)
            path = out
        item["aspect"] = upright.width / upright.height
    item.update(kind="image", file=path)
    return item


def video_poster(path, work):
    """First-frame poster via macOS Quick Look; a plain dark frame when that is unavailable."""
    out_dir = work / f"poster-{len(list(work.iterdir()))}"
    out_dir.mkdir()
    try:
        subprocess.run(["qlmanage", "-t", "-s", "1280", "-o", str(out_dir), str(path)],
                       capture_output=True, timeout=60)
        made = list(out_dir.glob("*.png"))
        if made:
            return made[0]
    except (OSError, subprocess.TimeoutExpired):
        pass
    warn(f"no poster frame for {path.name}; pass \"poster\" to use your own")
    poster = out_dir / "poster.png"
    im = Image.new("RGB", (1280, 720), (17, 31, 74))
    d = ImageDraw.Draw(im)
    d.polygon([(560, 260), (560, 460), (740, 360)], fill=(255, 255, 255))
    im.save(poster)
    return poster


def caption_height(text, width_emu, size=None):
    if not text:
        return 0
    size = size or STYLE["min_font_size"]
    runs = [Run(text)]
    return line_count(runs, width_emu, size) * line_height(runs, size) * EMU_PER_PT + 0.06 * EMU_PER_IN


def media_layout(items, box, rows=None, valign="top"):
    """Justified rows: each row spans the box width; rows are scaled together to fit its height."""
    x, y, w, h = box
    n = len(items)
    best = None
    for r in ([min(int(rows), n)] if rows else range(1, n + 1)):
        counts = [n // r + (1 if i < n % r else 0) for i in range(r)]
        rows_spec, idx = [], 0
        for c in counts:
            group = items[idx:idx + c]
            row_h = (w - GAP * (c - 1)) / sum(it["aspect"] for it in group)
            rows_spec.append((idx, c, row_h))
            idx += c
        fixed = GAP * (r - 1)
        for i0, c, row_h in rows_spec:
            fixed += max(caption_height(it.get("caption"), it["aspect"] * row_h) for it in items[i0:i0 + c])
        scale = min(1.0, (h - fixed) / sum(rh for _, _, rh in rows_spec))
        if scale <= 0:
            continue
        area = sum(it["aspect"] * (rh * scale) ** 2 for i0, c, rh in rows_spec for it in items[i0:i0 + c])
        if best is None or area > best[0]:
            best = (area, rows_spec, scale)
    if best is None:
        raise SystemExit("media do not fit: too many items for the space left on the slide")
    _, rows_spec, scale = best

    placed, cy = [], y
    for i0, c, row_h in rows_spec:
        rh = row_h * scale
        group = items[i0:i0 + c]
        row_w = sum(it["aspect"] * rh for it in group) + GAP * (c - 1)
        cx = x + (w - row_w) / 2
        cap = max(caption_height(it.get("caption"), it["aspect"] * rh) for it in group)
        for it in group:
            iw = it["aspect"] * rh
            placed.append((it, (cx, cy, iw, rh), cap))
            cx += iw + GAP
        cy += rh + cap + GAP
    used = cy - GAP - y
    shift = (h - used) / 2 if valign == "middle" else 0
    return [(it, (bx, by + shift, bw, bh), cap) for it, (bx, by, bw, bh), cap in placed], used


def add_media(slide, items, box, rows=None, valign="top"):
    placed, used = media_layout(items, box, rows, valign)
    for k, (it, (x, y, w, h), cap) in enumerate(placed, 1):
        x, y, w, h = int(x), int(y), int(w), int(h)
        if it["kind"] == "chart":
            add_chart(slide, it["chart"], (x, y, w, h), f"Chart {k}")
        elif it["kind"] == "video":
            shape = slide.shapes.add_movie(str(it["src"]), Emu(x), Emu(y), Emu(w), Emu(h),
                                           poster_frame_image=str(it["poster"]), mime_type=it["mime"])
            shape.name = f"Video {k}"
        else:
            shape = slide.shapes.add_picture(str(it["file"]), Emu(x), Emu(y), Emu(w), Emu(h))
            shape.name = f"Figure {k}"
        if it["kind"] != "chart":
            shape._element.xpath("./*[1]/p:cNvPr")[0].set("descr", it.get("caption") or it["src"].name)
        if it.get("caption"):
            tb = add_textbox(slide, (x - GAP / 2, y + h + 0.04 * EMU_PER_IN, w + GAP, cap), f"Caption {k}")
            write_paragraphs(tb.text_frame._txBody, [Para(inline_runs(it["caption"], color=SLATE), algn="ctr",
                                                          color=SLATE)], STYLE["min_font_size"])
    return used


CHART_COLORS = [NAVY_DARK, "5B8DB8", "C9A04A", ONGOING, "8A97B0"]
HIGHLIGHT = "C9A04A"   # the gold of the template's header rule


def add_chart(slide, spec, box, name):
    """Native PowerPoint chart (stays editable), text at the house minimum size."""
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
    from pptx.util import Pt

    kinds = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
             "line": XL_CHART_TYPE.LINE_MARKERS}
    kind = spec.get("type", "bar")
    if kind not in kinds:
        raise SystemExit(f"chart type {kind!r}: use bar, hbar or line")
    number_format = spec.get("number_format", "General")
    data = CategoryChartData(number_format=number_format)
    data.categories = [str(c) for c in spec["categories"]]
    for ser in spec["series"]:
        data.add_series(ser["name"], ser["values"])
    x, y, w, h = (int(v) for v in box)
    frame = slide.shapes.add_chart(kinds[kind], Emu(x), Emu(y), Emu(w), Emu(h), data)
    frame.name = name
    chart = frame.chart
    size = Pt(STYLE["min_font_size"])
    chart.font.size = size
    chart.font.name = "Calibri"
    chart.font.color.rgb = RGBColor.from_string(NAVY)
    chart.has_title = False
    multi = len(spec["series"]) > 1
    chart.has_legend = multi
    if multi:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False

    plot = chart.plots[0]
    if kind != "line":
        plot.gap_width = 70
        plot.overlap = -10 if multi else 0
    highlight = {str(c) for c in spec.get("highlight", [])}
    for i, series in enumerate(plot.series):
        color = RGBColor.from_string(spec["series"][i].get("color", CHART_COLORS[i % len(CHART_COLORS)]))
        if kind == "line":
            series.smooth = False
            series.format.line.color.rgb = color
            series.format.line.width = Pt(3)
            series.marker.size = 9
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
            series.marker.format.line.color.rgb = color
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
            if not multi:
                for j, cat in enumerate(data.categories):
                    if cat.label in highlight:
                        point = series.points[j].format.fill
                        point.solid()
                        point.fore_color.rgb = RGBColor.from_string(HIGHLIGHT)

    labels = spec.get("labels", kind != "line")
    value_axis = chart.value_axis
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = size
        dl.font.bold = True
        dl.font.color.rgb = RGBColor.from_string(NAVY)
        dl.number_format = number_format
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.ABOVE if kind == "line" else XL_LABEL_POSITION.OUTSIDE_END
        value_axis.visible = False          # the labels carry the numbers
        value_axis.has_major_gridlines = False
    else:
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(RULE)
        value_axis.format.line.fill.background()
        value_axis.tick_labels.font.size = size
    for key, attr in (("y_min", "minimum_scale"), ("y_max", "maximum_scale")):
        if spec.get(key) is not None:
            setattr(value_axis, attr, float(spec[key]))
    category_axis = chart.category_axis
    category_axis.tick_labels.font.size = size
    category_axis.format.line.color.rgb = RGBColor.from_string(RULE)
    category_axis.has_major_gridlines = False


# ---------------------------------------------------------------------------
# Tables


def cell_spec(value):
    if isinstance(value, dict):
        return {"text": str(value.get("text", "")), "bold": value.get("bold", False),
                "color": value.get("color"), "fill": value.get("fill")}
    return {"text": "" if value is None else str(value), "bold": False, "color": None, "fill": None}


CELL_PAD_X = int(0.08 * EMU_PER_IN)
CELL_PAD_Y = int(0.045 * EMU_PER_IN)


def plan_table(t, width, max_h, max_size=24):
    """Pick the largest font that fits the height without wrapping cells more than the minimum size does."""
    header = [cell_spec(v) for v in t.get("header", [])]
    rows = [[cell_spec(v) for v in r] for r in t["rows"]]
    grid = ([header] if header else []) + rows
    ncols = max(len(r) for r in grid)
    for r in grid:
        r.extend(cell_spec("") for _ in range(ncols - len(r)))

    def layout(size):
        table_w = width
        if t.get("col_widths"):
            ratios = [float(v) for v in t["col_widths"]]
        else:
            pad = 2 * CELL_PAD_X / EMU_PER_PT + 6
            natural = [max(METRICS.width(r[j]["text"], size, True) for r in grid) + pad for j in range(ncols)]
            cap = width / EMU_PER_PT * 0.45
            ratios = [min(max(v, 3 * size), cap) for v in natural]
            # A narrow table keeps a comfortable width instead of stretching across the slide.
            table_w = min(width, int(sum(ratios) * 1.5 * EMU_PER_PT))
        total = sum(ratios)
        widths = [int(table_w * v / total) for v in ratios]
        widths[-1] = table_w - sum(widths[:-1])
        heights, wrapped = [], 0
        for i, r in enumerate(grid):
            bold_row = bool(header) and i == 0
            lines = [line_count(inline_runs(c["text"], bold=bold_row or c["bold"]), widths[j] - 2 * CELL_PAD_X, size)
                     for j, c in enumerate(r)]
            wrapped += sum(lines)
            lh = max(line_height(inline_runs(c["text"]), size) for c in r)
            heights.append(int(max(lines) * lh * EMU_PER_PT + 2 * CELL_PAD_Y))
        return {"widths": widths, "heights": heights, "size": size, "lines": wrapped}

    floor = layout(STYLE["min_font_size"])
    if t.get("font_size"):
        plan = layout(float(t["font_size"]))
        if plan["lines"] > floor["lines"]:
            warn(f"table at {plan['size']:g}pt wraps cells that fit on one line at smaller sizes — "
                 "drop font_size or shorten the cells")
    else:
        plans = [layout(s) for s in font_sizes(24) if s <= max(max_size, STYLE["min_font_size"])]
        fits = [p for p in plans if sum(p["heights"]) <= max_h]
        plan = next((p for p in fits if p["lines"] <= floor["lines"]), fits[-1] if fits else floor)
    if sum(plan["heights"]) > max_h:
        warn(f"table with {len(rows)} rows is taller than the space left at {plan['size']:g}pt — split it")
    plan.update(grid=grid, header=bool(header), align=t.get("align"))
    return plan


def add_table(slide, plan, x, y):
    grid, widths, heights = plan["grid"], plan["widths"], plan["heights"]
    shape = slide.shapes.add_table(len(grid), len(widths), Emu(x), Emu(y), Emu(sum(widths)), Emu(sum(heights)))
    shape.name = "Table"
    table = shape.table
    table.first_row = plan["header"]
    table.horz_banding = False
    for j, w in enumerate(widths):
        table.columns[j].width = Emu(w)
    for i, h in enumerate(heights):
        table.rows[i].height = Emu(h)
    for i, r in enumerate(grid):
        is_head = plan["header"] and i == 0
        body_i = i - 1 if plan["header"] else i
        for j, c in enumerate(r):
            cell = table.cell(i, j)
            if plan["align"]:
                algn = {"left": "l", "center": "ctr", "right": "r"}[plan["align"][j]]
            else:
                algn = "ctr" if is_head or j > 0 else "l"
            color = "FFFFFF" if is_head else (c["color"] or NAVY)
            runs = inline_runs(c["text"], bold=is_head or c["bold"], color=color)
            write_paragraphs(cell.text_frame._txBody, [Para(runs, algn=algn, color=color)], plan["size"])
            fill = c["fill"] or (NAVY_DARK if is_head else ("FFFFFF" if body_i % 2 == 0 else STRIPE))
            style_cell(cell, fill)


def style_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        tc_pr.remove(child)
    tc_pr.set("marL", str(CELL_PAD_X))
    tc_pr.set("marR", str(CELL_PAD_X))
    tc_pr.set("marT", str(CELL_PAD_Y))
    tc_pr.set("marB", str(CELL_PAD_Y))
    tc_pr.set("anchor", "ctr")
    for tag in ("lnL", "lnR", "lnT", "lnB"):
        tc_pr.append(parse_xml(
            f'<a:{tag} {nsdecls("a")} w="6350" cap="flat" cmpd="sng" algn="ctr">'
            f'<a:solidFill><a:srgbClr val="{RULE}"/></a:solidFill><a:prstDash val="solid"/>'
            f'<a:round/><a:headEnd type="none" w="med" len="med"/><a:tailEnd type="none" w="med" len="med"/>'
            f'</a:{tag}>'))
    tc_pr.append(parse_xml(f'<a:solidFill {nsdecls("a")}><a:srgbClr val="{fill}"/></a:solidFill>'))


# ---------------------------------------------------------------------------
# Content slide


def fill_content(slide, s, base_dir, work):
    set_title(slide, s["title"])
    x0, x1 = MARGIN_X, SLIDE_W - MARGIN_X
    y0, y1 = CONTENT_TOP, CONTENT_BOTTOM
    label = f"slide '{s['title']}'"

    if s.get("subtitle"):
        paras = [Para(inline_runs(s["subtitle"], bold=True, color=NAVY_DARK), color=NAVY_DARK)]
        used, _ = place_text(slide, paras, (x0, y0, x1 - x0, 1.0 * EMU_PER_IN), font_sizes(24), "Subtitle",
                             f"{label} subtitle")
        y0 += used + GAP
    bullets = bullet_paras(s.get("bullets", []))
    media = [prepare_media(m, base_dir, work) for m in s.get("media", [])]
    table = s.get("table")
    layout = s.get("layout", "auto")
    if layout == "fit":
        # Figure-led slide: a wide figure gets the full width under the text, a squarer one sits beside it.
        layout = "stack" if media and all(m["aspect"] >= WIDE_FIGURE for m in media) else "auto"
    if layout == "auto":
        layout = "text-left" if bullets and media and not table else "stack"

    if layout in ("text-left", "text-right") and bullets:
        tw = int((x1 - x0) * float(s.get("text_width", 0.38)))
        if layout == "text-left":
            text_box, rest = (x0, y0, tw, y1 - y0), (x0 + tw + COL_GAP, y0, x1, y1)
        else:
            text_box, rest = (x1 - tw, y0, tw, y1 - y0), (x0, y0, x1 - tw - COL_GAP, y1)
        place_text(slide, bullets, text_box, font_sizes(24), "Bullets", f"{label} bullets")
        bullets = []
        x0, y0, x1, y1 = rest

    table_max = 24
    if bullets:
        alone = not media and not table
        room = (y1 - y0) if alone else (y1 - y0) * 0.45
        sizes = font_sizes(28) if alone else font_sizes(24)
        used, table_max = place_text(slide, bullets, (x0, y0, x1 - x0, room), sizes, "Bullets", f"{label} bullets")
        y0 += used + GAP

    if table:
        room = (y1 - y0) if not media else (y1 - y0) * 0.55
        plan = plan_table(table, x1 - x0, room, table_max)
        add_table(slide, plan, x0 + (x1 - x0 - sum(plan["widths"])) // 2, y0)
        y0 += sum(plan["heights"]) + GAP

    if media:
        if y1 - y0 < 0.8 * EMU_PER_IN:
            raise SystemExit(f"{label}: no room left for media — move text or the table to another slide")
        only_media = not s.get("bullets") and not table
        add_media(slide, media, (x0, y0, x1 - x0, y1 - y0), s.get("media_rows"),
                  "middle" if only_media else "top")

    if s.get("valign") == "middle":
        center_vertically([sh for sh in slide.shapes if sh.shape_id != slide.shapes.title.shape_id])

    if s.get("footnote"):
        paras = [Para(inline_runs(s["footnote"], color=SLATE), color=SLATE)]
        place_text(slide, paras, (MARGIN_X, FOOTNOTE_TOP, LOGO_LEFT - MARGIN_X - GAP, FOOTNOTE_BOTTOM - FOOTNOTE_TOP),
                   [12, 11, 10], "Footnote", f"{label} footnote")


def center_vertically(shapes):
    """Move everything below the header down so the block sits in the middle of the content area."""
    shapes = [sh for sh in shapes if sh.top >= CONTENT_TOP - EMU_PER_PT and sh.top < CONTENT_BOTTOM]
    if not shapes:
        return
    top = min(sh.top for sh in shapes)
    bottom = max(sh.top + sh.height for sh in shapes)
    shift = CONTENT_TOP + (CONTENT_BOTTOM - CONTENT_TOP - (bottom - top)) // 2 - top
    if shift > 0:
        for sh in shapes:
            sh.top = Emu(sh.top + shift)


# ---------------------------------------------------------------------------
# Spec handling


def load_config():
    if CONFIG.exists():
        return json.loads(CONFIG.read_text(encoding="utf-8"))
    return {}


def parse_date(date):
    """'2026.10.29', '2026-10-29', '20261029' or '261029' -> datetime.date (None if unreadable)."""
    digits = re.findall(r"\d+", str(date))
    if len(digits) == 1 and len(digits[0]) in (6, 8):
        v = digits[0]
        digits = [v[:-4], v[-4:-2], v[-2:]]
    if len(digits) < 3:
        return None
    y, m, d = (int(v) for v in digits[:3])
    try:
        return datetime.date(y + 2000 if y < 100 else y, m, d)
    except ValueError:
        return None


def slide_texts(s):
    texts = [s.get("title"), s.get("subtitle"), s.get("name")] + list(s.get("extra_lines", []))
    texts += [normalize_item(b)["text"] for b in s.get("bullets", [])]
    for proj in s.get("projects", []):
        texts += [proj.get("name")] + [normalize_item(i)["text"] for i in proj.get("items", [])]
    table = s.get("table") or {}
    for row in [table.get("header", [])] + list(table.get("rows", [])):
        texts += [cell_spec(c)["text"] for c in row]
    for m in s.get("media", []):
        if isinstance(m, dict):
            chart = m.get("chart") or {}
            texts += [m.get("caption")] + [str(c) for c in chart.get("categories", [])]
            texts += [ser.get("name") for ser in chart.get("series", [])]
    return [t for t in texts if t]


def words(text):
    return len(text.replace("**", "").split())


def lint(s):
    """House-style checks: English only, a few short key sentences, a visual on every content slide."""
    where = f"slide '{s.get('title') or 'Group Meeting'}'"
    if STYLE["language"] == "en":
        foreign = next((t for t in slide_texts(s) if CJK_RE.search(t)), None)
        if foreign:
            warn(f"{where}: non-English text {foreign!r} — the house style is all English")
    limit = STYLE["max_words_per_bullet"]
    if s["type"] == "content":
        bullets = [normalize_item(b)["text"] for b in s.get("bullets", [])]
        if len(bullets) > STYLE["max_bullets_per_slide"]:
            warn(f"{where}: {len(bullets)} bullets — keep to {STYLE['max_bullets_per_slide']} key sentences "
                 "or split the slide")
        long = next((b for b in bullets if words(b) > limit), None)
        if long:
            warn(f"{where}: bullet over {limit} words — {long!r}")
        if not s.get("media") and not s.get("text_only"):
            warn(f"{where}: no figure, demo video or chart — add a visual (or set \"text_only\": true)")
    for proj in s.get("projects", []):
        long = next((normalize_item(i)["text"] for i in proj.get("items", [])
                     if words(normalize_item(i)["text"]) > limit), None)
        if long:
            warn(f"{where}: item over {limit} words — {long!r}")


def as_list(value):
    if value is None or value == "":
        return []
    return [value] if isinstance(value, str) else list(value)


def paper_slides(s):
    """One paper -> two content slides: pipeline + problem/approach, then novelty | limitations."""
    name = s.get("title") or s.get("name")
    if not name:
        raise SystemExit("every paper slide needs a short \"title\" (the paper's name)")
    where = f"paper '{name}'"
    problem, approach = s.get("problem"), s.get("approach")
    novelty, limits = as_list(s.get("novelty")), as_list(s.get("limitations"))
    needed = (("the original pipeline figure", s.get("pipeline")), ("the problem it solves", problem),
              ("how it solves it", approach), ("the novelty", novelty), ("the limitations", limits))
    missing = [label for label, value in needed if not value]
    if missing:
        warn(f"{where}: missing {', '.join(missing)} — ask the user for them, don't fill them in yourself")
    notes = as_list(s.get("notes"))

    bullets = [f"**{label}:** {text}" for label, text in (("Problem", problem), ("Approach", approach)) if text]
    first = {"type": "content", "title": f"Paper reading: {name}", "bullets": bullets, "layout": "fit",
             "text_only": not s.get("pipeline")}
    if s.get("full_title"):
        first["subtitle"] = s["full_title"]
    if s.get("pipeline"):
        first["media"] = [{"path": s["pipeline"], "caption": s.get("pipeline_caption", "")}]
    source = [t for t in (s.get("venue"), "pipeline figure from the paper" if s.get("pipeline") else None) if t]
    if source:
        first["footnote"] = " · ".join(source)
    first.update({k: s[k] for k in ("layout", "text_width") if k in s})
    if notes:
        first["notes"] = notes[0]
    out = [first]

    columns = [(head, items) for head, items in (("Novelty", novelty), ("Limitations", limits)) if items]
    if columns:
        for head, items in columns:
            if len(items) > 2:
                warn(f"{where}: {len(items)} {head.lower()} — keep the two that matter most")
        long = next((t for _, items in columns for t in items if words(t) > STYLE["max_words_per_bullet"]), None)
        if long:
            warn(f"{where}: point over {STYLE['max_words_per_bullet']} words — {long!r}")
        count = max(len(items) for _, items in columns)
        rows = [[items[i] if i < len(items) else "" for _, items in columns] for i in range(count)]
        second = {"type": "content", "title": f"{name}: novelty and limits", "text_only": True, "valign": "middle",
                  "table": {"header": [head for head, _ in columns], "rows": rows,
                            "col_widths": [1] * len(columns), "align": ["left"] * len(columns)}}
        if s.get("takeaway"):
            second["subtitle"] = s["takeaway"]
        if len(notes) > 1:
            second["notes"] = notes[1]
        out.append(second)
    return out


NUMBER_RE = re.compile(r"-?\d+(?:\.\d+)?")


def bold_best(table):
    """Rows of a results table with the best value of each column bold (from ↓/↑ in the header, or `better`)."""
    header = [cell_spec(c)["text"] for c in table.get("header", [])]
    rows = [list(r) for r in table["rows"]]
    better = table.get("better") or []
    for j in range(max(len(r) for r in rows)):
        rule = better[j] if j < len(better) else None
        if rule is None and j < len(header):
            rule = "lower" if "↓" in header[j] else "higher" if "↑" in header[j] else None
        if rule not in ("lower", "higher"):
            continue
        found = []
        for i, r in enumerate(rows):
            hit = NUMBER_RE.search(cell_spec(r[j])["text"]) if j < len(r) else None
            if hit:
                found.append((float(hit.group()), i))
        if not found:
            continue
        best = (min if rule == "lower" else max)(v for v, _ in found)
        for v, i in found:
            if v == best:
                rows[i][j] = dict(cell_spec(rows[i][j]), bold=True)
    return rows


def compare_slides(s):
    """One theme's papers -> a methods table slide, then a results table slide. Tables only, no figures."""
    theme = s.get("theme") or s.get("title")
    if not theme:
        raise SystemExit("every compare slide needs a \"theme\" (the group of papers it compares)")
    where = f"comparison '{theme}'"
    methods, results = s.get("methods"), s.get("results")
    if not methods or not methods.get("rows"):
        raise SystemExit(f"{where}: needs a \"methods\" table — compare how the papers work before their results")
    if not results or not results.get("rows"):
        warn(f"{where}: no results table — ask the user for the numbers on a shared benchmark")
    notes = as_list(s.get("notes"))

    out = []
    for kind, t in (("methods", methods), ("results", results)):
        if not t or not t.get("rows"):
            continue
        ncols = max(len(r) for r in ([t.get("header", [])] + list(t["rows"])))
        if len(t["rows"]) > STYLE["max_compare_papers"]:
            warn(f"{where} {kind}: {len(t['rows'])} papers — keep {STYLE['max_compare_papers']} or fewer per "
                 "theme, or split the theme")
        if ncols > STYLE["max_compare_columns"]:
            warn(f"{where} {kind}: {ncols} columns — keep {STYLE['max_compare_columns']} or fewer "
                 "(paper name + the few dimensions that matter)")
        table = {k: t[k] for k in ("header", "col_widths", "font_size") if k in t}
        table["rows"] = bold_best(t) if kind == "results" else t["rows"]
        table["align"] = t.get("align") or (["left"] * ncols if kind == "methods" else ["left"] + ["center"] * (ncols - 1))
        slide = {"type": "content", "title": f"Compare: {theme} — {kind}", "table": table, "text_only": True}
        if t.get("takeaway"):
            slide["subtitle"] = t["takeaway"]
        if kind == "results":
            if t.get("note"):
                slide["footnote"] = t["note"]
            else:
                warn(f"{where} results: no \"note\" — say which benchmark and setting the numbers come from")
        if len(notes) > len(out):
            slide["notes"] = notes[len(out)]
        out.append(slide)
    return out


def expand_slides(spec, cfg):
    slides = []
    for s in spec["slides"]:
        kind = s.get("type", "content")
        slides.extend(paper_slides(s) if kind == "paper" else compare_slides(s) if kind == "compare" else [dict(s)])
    if not any(s.get("type") == "cover" for s in slides) and spec.get("cover", True) is not False:
        slides.insert(0, {"type": "cover"})
    planned = []
    for s in slides:
        kind = s.get("type", "content")
        if kind not in BASE_SLIDE:
            raise SystemExit(f"unknown slide type {kind!r}; use cover, done, content, paper, compare or todo")
        s["type"] = kind
        if kind == "cover":
            day = parse_date(s["date"]) if s.get("date") else datetime.date.today()
            if day:
                s["date"] = day.strftime("%Y.%m.%d")
                s["_day"] = day
            else:
                warn(f"cover date {s['date']!r} is not a date; used as written")
            s.setdefault("name", cfg.get("presenter", ""))
            if not s["name"]:
                warn("cover has no presenter name — set \"name\" in the spec or \"presenter\" in config.json")
            lint(s)
            planned.append(s)
        elif kind in ("done", "todo"):
            s.setdefault("title", cfg.get(f"{kind}_title") or ("Done last two weeks" if kind == "done" else "Todo"))
            lint(s)
            planned.extend(plan_progress(s, kind))
        else:
            if not s.get("title"):
                raise SystemExit("every content slide needs a title")
            lint(s)
            planned.append(s)
    return planned


def find_pptx_skill(given):
    candidates = []
    if given:
        candidates.append(Path(given).expanduser())
    if os.environ.get("PPTX_SKILL_DIR"):
        candidates.append(Path(os.environ["PPTX_SKILL_DIR"]).expanduser())
    claude = Path.home() / ".claude"
    for pattern in ("skills/pptx", "skills/*/pptx", "skills/*/*/pptx", "plugins/**/skills/pptx"):
        candidates.extend(sorted(claude.glob(pattern)))
    for c in candidates:
        if (c / "scripts" / "add_slide.py").exists():
            return c.resolve()
    raise SystemExit("pptx skill not found. Load it with the Skill tool and pass its base directory "
                     "as --pptx-skill (or set PPTX_SKILL_DIR).")


def run(cmd):
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        raise SystemExit(f"{' '.join(map(str, cmd))} failed:\n{res.stdout}{res.stderr}")
    return res.stdout


def build_structure(kinds, scripts, out_path):
    """Clone template sample slides in spec order with the pptx skill's scripts."""
    sys.path.insert(0, str(scripts))
    from office.helpers import rezip, safe_extract

    with tempfile.TemporaryDirectory(prefix="gm_struct_") as tmp:
        unpacked = Path(tmp) / "deck"
        with zipfile.ZipFile(TEMPLATE) as zf:
            safe_extract(zf, unpacked)
        for kind in kinds:
            run([sys.executable, str(scripts / "add_slide.py"), str(unpacked), BASE_SLIDE[kind]])

        rels = (unpacked / "ppt" / "_rels" / "presentation.xml.rels").read_text(encoding="utf-8")
        originals = set()
        for rel in re.findall(r"<Relationship\b[^>]*>", rels):
            target = re.search(r'Target="([^"]+)"', rel).group(1)
            if target.rsplit("/", 1)[-1] in TEMPLATE_SLIDES and "/slides/" in "/" + target:
                originals.add(re.search(r'Id="([^"]+)"', rel).group(1))
        pres_path = unpacked / "ppt" / "presentation.xml"
        pres = pres_path.read_text(encoding="utf-8")
        pres = re.sub(r'<p:sldId\b[^>]*r:id="([^"]+)"[^>]*/>',
                      lambda m: "" if m.group(1) in originals else m.group(0), pres)
        pres_path.write_text(pres, encoding="utf-8")

        run([sys.executable, str(scripts / "clean.py"), str(unpacked)])
        rezip(unpacked, out_path)


def main():
    global METRICS
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", help="deck spec (JSON)")
    ap.add_argument("-o", "--output", help="output .pptx (default: spec 'output', else <output_root>/<YYMMDD>GroupMeeting/)")
    ap.add_argument("--pptx-skill", help="base directory of the pptx skill (auto-detected under ~/.claude)")
    ap.add_argument("--no-validate", action="store_true", help="skip the pptx skill's validate.py")
    args = ap.parse_args()

    METRICS = Metrics()
    if "latin" not in METRICS.fonts:
        warn("Calibri not found locally — text fitting uses estimated widths")

    spec_path = Path(args.spec).resolve()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    cfg = load_config()
    STYLE.update(cfg.get("style", {}))
    planned = expand_slides(spec, cfg)
    cover = next((s for s in planned if s["type"] == "cover"), None)
    code = ((cover or {}).get("_day") or datetime.date.today()).strftime("%y%m%d")

    if args.output:
        out = Path(os.path.expanduser(args.output)).resolve()
    elif spec.get("output"):
        out = Path(os.path.expanduser(spec["output"]))
        out = out if out.is_absolute() else (spec_path.parent / out).resolve()
    else:
        root = Path(os.path.expanduser(cfg.get("output_root", "~/Documents/PairsLab")))
        out = root / f"{code}GroupMeeting" / f"{code}GroupMeeting.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)

    pptx_skill = find_pptx_skill(args.pptx_skill)
    scripts = pptx_skill / "scripts"
    build_structure([s["type"] for s in planned], scripts, out)

    prs = Presentation(str(out))
    with tempfile.TemporaryDirectory(prefix="gm_media_") as tmp:
        work = Path(tmp)
        for slide, s in zip(prs.slides, planned):
            if s["type"] == "cover":
                fill_cover(slide, s)
            elif s["type"] in ("done", "todo"):
                fill_progress(slide, s, s["type"])
            else:
                fill_content(slide, s, spec_path.parent, work)
            if s.get("notes"):
                slide.notes_slide.notes_text_frame.text = s["notes"]
        prs.save(str(out))

    print(f"Wrote {out} ({len(planned)} slides)")
    for i, s in enumerate(planned, 1):
        print(f"  {i:2d}. [{s['type']}] {s.get('title') or 'Group Meeting'}")
    for w in WARNINGS:
        print(f"WARNING: {w}")

    if not args.no_validate:
        res = subprocess.run([sys.executable, str(scripts / "office" / "validate.py"), str(out),
                              "--original", str(TEMPLATE)], capture_output=True, text=True)
        print((res.stdout + res.stderr).strip())
        if res.returncode != 0:
            sys.exit(1)


if __name__ == "__main__":
    main()
